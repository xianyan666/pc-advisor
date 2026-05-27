from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


OUT = Path("项目报告") / "电脑硬件分析推荐系统_答辩展示_修改版.pptx"

BG = RGBColor(239, 247, 255)
BG2 = RGBColor(226, 241, 255)
WHITE = RGBColor(255, 255, 255)
BLUE = RGBColor(50, 128, 220)
BLUE_DARK = RGBColor(25, 74, 145)
BLUE_SOFT = RGBColor(213, 232, 255)
CYAN = RGBColor(71, 177, 230)
TEXT = RGBColor(35, 48, 68)
MUTED = RGBColor(94, 111, 132)
LINE = RGBColor(185, 211, 240)
GREEN = RGBColor(38, 159, 116)


def rgb(hex_color):
    hex_color = hex_color.lstrip("#")
    return RGBColor(int(hex_color[0:2], 16), int(hex_color[2:4], 16), int(hex_color[4:6], 16))


def run_style(run, size=14, bold=False, color=TEXT):
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def text(slide, x, y, w, h, value, size=14, bold=False, color=TEXT, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = Inches(0.04)
    tf.margin_right = Inches(0.04)
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = value
    run_style(r, size, bold, color)
    return box


def bullets(slide, x, y, w, h, items, size=11.2, color=TEXT, gap=4):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.03)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        r = p.add_run()
        r.text = f"• {item}"
        run_style(r, size, False, color)
    return box


def background(slide, prs):
    rect = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    rect.fill.solid()
    rect.fill.fore_color.rgb = BG
    rect.line.fill.background()
    blob = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, Inches(0.28))
    blob.fill.solid()
    blob.fill.fore_color.rgb = BLUE
    blob.line.fill.background()
    side = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(12.84), Inches(0), Inches(0.18), prs.slide_height)
    side.fill.solid()
    side.fill.fore_color.rgb = BLUE_SOFT
    side.line.fill.background()


def title(slide, value, subtitle=None):
    text(slide, 0.62, 0.47, 9.8, 0.38, value, 23, True, BLUE_DARK)
    if subtitle:
        text(slide, 0.64, 0.88, 11.2, 0.27, subtitle, 10.5, False, MUTED)
    line = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.64), Inches(1.23), Inches(1.25), Inches(0.05))
    line.fill.solid()
    line.fill.fore_color.rgb = CYAN
    line.line.fill.background()


def footer(slide, page):
    text(slide, 0.64, 7.12, 5.5, 0.22, "电脑硬件分析推荐系统 | 软件工程综合课设答辩", 8.5, False, MUTED)
    text(slide, 12.25, 7.12, 0.45, 0.22, str(page), 8.5, True, BLUE_DARK, PP_ALIGN.RIGHT)


def card(slide, x, y, w, h, heading, items, accent=BLUE, heading_size=13.2, body_size=10.5):
    s = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = WHITE
    s.line.color.rgb = LINE
    s.line.width = Pt(1.1)
    top = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.1))
    top.fill.solid()
    top.fill.fore_color.rgb = accent
    top.line.fill.background()
    text(slide, x + 0.22, y + 0.22, w - 0.4, 0.28, heading, heading_size, True, BLUE_DARK)
    if isinstance(items, str):
        items = [items]
    bullets(slide, x + 0.22, y + 0.62, w - 0.4, h - 0.72, items, body_size)
    return s


def pill(slide, x, y, w, label, color=BLUE):
    p = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.34))
    p.fill.solid()
    p.fill.fore_color.rgb = color
    p.line.fill.background()
    text(slide, x, y + 0.07, w, 0.18, label, 9.2, True, WHITE, PP_ALIGN.CENTER)


def flow_box(slide, x, y, w, title_value, body):
    s = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(1.05))
    s.fill.solid()
    s.fill.fore_color.rgb = WHITE
    s.line.color.rgb = LINE
    text(slide, x + 0.1, y + 0.16, w - 0.2, 0.22, title_value, 11.5, True, BLUE_DARK, PP_ALIGN.CENTER)
    text(slide, x + 0.12, y + 0.52, w - 0.24, 0.25, body, 9.2, False, MUTED, PP_ALIGN.CENTER)


def arrow(slide, x1, y1, x2, y2):
    line = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line.line.color.rgb = BLUE
    line.line.width = Pt(1.4)
    line.line.end_arrowhead = True


def section(prs, page, heading, subheading, tags):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = BLUE_DARK
    bg.line.fill.background()
    soft = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(6.5), prs.slide_width, Inches(1))
    soft.fill.solid()
    soft.fill.fore_color.rgb = rgb("#2f86de")
    soft.line.fill.background()
    text(slide, 0.9, 2.35, 11.5, 0.6, heading, 31, True, WHITE, PP_ALIGN.CENTER)
    text(slide, 1.25, 3.12, 10.8, 0.35, subheading, 13.5, False, RGBColor(220, 235, 255), PP_ALIGN.CENTER)
    start = 3.9
    for i, tag in enumerate(tags):
        pill(slide, start + i * 1.85, 4.02, 1.5, tag, rgb("#5fb3f3"))
    footer(slide, page)
    return slide


prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

page = 1

# Cover
slide = prs.slides.add_slide(prs.slide_layouts[6])
background(slide, prs)
panel = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.85), Inches(1.05), Inches(11.65), Inches(5.55))
panel.fill.solid()
panel.fill.fore_color.rgb = WHITE
panel.line.color.rgb = LINE
panel.line.width = Pt(1.2)
text(slide, 1.35, 1.75, 10.5, 0.68, "电脑硬件分析推荐系统", 34, True, BLUE_DARK, PP_ALIGN.CENTER)
text(slide, 1.65, 2.55, 10.0, 0.35, "软件工程综合课设答辩 | 淡蓝白简约版", 15, False, MUTED, PP_ALIGN.CENTER)
for i, label in enumerate(["后端 + 数据库", "前端 + 项目管理", "APP 端 + 报告撰写"]):
    pill(slide, 2.05 + i * 3.15, 3.35, 2.45, label, [BLUE, CYAN, GREEN][i])
text(slide, 1.75, 4.45, 9.9, 0.7, "系统目标：围绕电脑硬件选购场景，完成硬件信息管理、排行比较、个性化推荐、评测社区、管理员审核与 Android 移动端访问。", 16, True, BLUE_DARK, PP_ALIGN.CENTER)
footer(slide, page)
page += 1

# System overview
slide = prs.slides.add_slide(prs.slide_layouts[6])
background(slide, prs)
title(slide, "01 系统完整描述", "先说明业务闭环，再展开三位同学负责的实现部分")
card(slide, 0.68, 1.52, 3.7, 4.85, "用户侧做什么", [
    "注册登录后进入系统，浏览 CPU、显卡、主板等硬件",
    "查看硬件详情、价格、品牌、型号和图片",
    "按榜单比较硬件性价比，按需求生成推荐配置",
    "收藏硬件，发布评测和评论，形成用户反馈内容",
], BLUE)
card(slide, 4.82, 1.52, 3.7, 4.85, "管理侧做什么", [
    "管理员登录后进入独立后台",
    "维护硬件基础数据和不同类型的专属参数",
    "审核用户发布的评测和评论，保证内容质量",
    "查看系统统计数据，掌握硬件和社区内容情况",
], CYAN)
card(slide, 8.95, 1.52, 3.7, 4.85, "系统逻辑是什么", [
    "前端负责交互和页面展示，后端负责业务规则和数据接口",
    "MySQL 保存用户、硬件、评测、评论、收藏和媒体资源",
    "推荐算法先筛选可用候选，再按偏好打分排序",
    "APP 复用 Web 页面，通过局域网访问同一套后端服务",
], GREEN)
footer(slide, page)
page += 1

# Architecture
slide = prs.slides.add_slide(prs.slide_layouts[6])
background(slide, prs)
title(slide, "02 总体架构与数据流", "一套后端和数据库同时支撑 Web 端、管理端和 Android 端")
flow_box(slide, 0.8, 2.0, 2.1, "用户 Web 端", "Vue 页面交互")
flow_box(slide, 0.8, 3.45, 2.1, "Android APP", "WebView 加载本地资源")
flow_box(slide, 3.8, 2.72, 2.25, "Axios / API 请求", "携带 JWT Token")
flow_box(slide, 6.75, 2.72, 2.2, "Spring Boot 后端", "Controller / Service / Mapper")
flow_box(slide, 9.75, 2.72, 2.25, "MySQL 数据库", "pcadvisor 数据库")
arrow(slide, 2.9, 2.52, 3.8, 3.02)
arrow(slide, 2.9, 3.95, 3.8, 3.22)
arrow(slide, 6.05, 3.25, 6.75, 3.25)
arrow(slide, 8.95, 3.25, 9.75, 3.25)
card(slide, 0.95, 5.25, 11.35, 1.0, "架构说明", [
    "普通用户、管理员和 APP 使用的是同一套后端接口，因此数据一致；前端只负责展示和交互，核心业务规则集中在后端。",
])
footer(slide, page)
page += 1

# Backend section
section(prs, page, "后端 + 数据库", "负责接口、业务逻辑、权限控制、数据库模型和推荐排行算法", ["架构", "数据", "算法"])
page += 1

slide = prs.slides.add_slide(prs.slide_layouts[6])
background(slide, prs)
title(slide, "后端做了什么", "按 Controller、Service、Mapper 分层实现业务")
card(slide, 0.7, 1.45, 3.75, 4.95, "接口层 Controller", [
    "AuthController：登录、注册、登出、用户名/邮箱检查",
    "HelloController：硬件列表、评测列表等公开接口",
    "RecommendationController：接收需求并返回推荐结果",
    "RankingController：生成 CPU、显卡、主板榜单",
    "AdminController：硬件管理、评测审核、评论审核、统计数据",
], BLUE)
card(slide, 4.8, 1.45, 3.75, 4.95, "业务层 Service", [
    "AuthService / UserService：用户认证和资料处理",
    "RecommendationService：配置推荐核心逻辑",
    "安全工具类：JWT 生成、解析和校验",
    "异常处理：统一返回错误码和错误信息",
], CYAN)
card(slide, 8.9, 1.45, 3.75, 4.95, "数据访问层 Mapper", [
    "基于 MyBatis-Plus 操作 MySQL",
    "按实体类映射用户、硬件、评测、评论等表",
    "使用 QueryWrapper 构造条件查询",
    "降低手写 SQL 成本，提高开发效率",
], GREEN)
footer(slide, page)
page += 1

slide = prs.slides.add_slide(prs.slide_layouts[6])
background(slide, prs)
title(slide, "数据库做了什么", "核心是把硬件领域对象抽象为可维护的数据模型")
card(slide, 0.7, 1.4, 3.75, 4.98, "用户与权限", [
    "sys_user 保存登录账号、密码和角色",
    "OrdinaryUser 保存普通用户扩展资料",
    "Administrator 保存管理员部门和等级",
    "配合 JWT 和 Spring Security 实现分权访问",
], BLUE)
card(slide, 4.8, 1.4, 3.75, 4.98, "硬件数据", [
    "Hardware_Inf 保存硬件通用字段：名称、品牌、型号、价格、审核状态",
    "CPU_Inf 保存核心数、线程数、频率、接口等",
    "GraphicsCard_Inf 保存显存、功耗、长度等",
    "Motherboard_Inf 保存板型、内存插槽、M.2 插槽等",
], CYAN)
card(slide, 8.9, 1.4, 3.75, 4.98, "社区与资源", [
    "Evaluation 保存评测标题、体验、优缺点和审核状态",
    "Comment 保存评测评论及审核状态",
    "Collect 保存用户收藏关系",
    "Media_Inf 保存图片路径，并关联硬件、评测或评论",
], GREEN)
text(slide, 0.9, 6.55, 11.5, 0.32, "逻辑：公共字段放主表，不同硬件的专属参数放详情表，避免一张大表字段冗余，也方便后续扩展内存、硬盘、电源等类型。", 12.2, True, BLUE_DARK)
footer(slide, page)
page += 1

slide = prs.slides.add_slide(prs.slide_layouts[6])
background(slide, prs)
title(slide, "推荐算法与安全逻辑", "算法解决“怎么推荐”，安全解决“谁能访问”")
card(slide, 0.72, 1.5, 5.75, 4.8, "推荐逻辑", [
    "输入：用途、预算、品牌偏好、核心数、接口、板型等条件",
    "第一步：从审核通过的硬件中取候选数据",
    "第二步：按预算、用途和兼容性进行硬性筛选",
    "第三步：按品牌、系列、板型等软性条件打分",
    "第四步：按得分和价格排序，返回 Top 推荐组合",
], BLUE)
card(slide, 6.85, 1.5, 5.75, 4.8, "安全逻辑", [
    "用户登录后生成 JWT Token，前端请求自动携带",
    "后端过滤器解析 Token 并设置当前用户身份",
    "管理员接口限制为 ADMIN 角色访问",
    "上传文件限制类型和大小，避免非法资源进入系统",
    "MyBatis-Plus 参数化查询降低 SQL 注入风险",
], CYAN)
footer(slide, page)
page += 1

# Front section
section(prs, page, "前端 + 项目管理", "负责页面交互、路由状态、接口联调、进度管理和功能验收", ["页面", "联调", "进度"])
page += 1

slide = prs.slides.add_slide(prs.slide_layouts[6])
background(slide, prs)
title(slide, "前端做了什么", "围绕普通用户和管理员两条使用路径组织页面")
card(slide, 0.7, 1.38, 3.75, 5.05, "普通用户页面", [
    "LoginView / RegisterView：登录注册与表单校验",
    "HomeView：首页推荐入口、分类入口和评测入口",
    "HardwareView / DetailView：硬件列表、筛选、详情、收藏",
    "RankingView / RecommendView：榜单查看和需求推荐",
    "EvaluationsView / UserView：评测社区和个人中心",
], BLUE)
card(slide, 4.8, 1.38, 3.75, 5.05, "管理员页面", [
    "AdminLayout：后台整体布局和菜单",
    "AdminDashboard：硬件、评测、评论等统计数据",
    "AdminHardware：硬件增删改查和分类参数维护",
    "AdminAudit：评测与评论的通过、驳回、删除",
], CYAN)
card(slide, 8.9, 1.38, 3.75, 5.05, "交互体验", [
    "使用 Vue Router 管理页面跳转和权限路由",
    "使用 Pinia 保存登录状态和用户角色",
    "使用 Axios 封装 API 请求和错误处理",
    "图片、图表、弹窗和响应式布局提升展示效果",
], GREEN)
footer(slide, page)
page += 1

slide = prs.slides.add_slide(prs.slide_layouts[6])
background(slide, prs)
title(slide, "前后端联调逻辑", "前端负责把用户操作转成 API 请求，后端返回结构化结果")
flow_box(slide, 0.85, 2.0, 2.0, "用户操作", "点击、筛选、提交表单")
flow_box(slide, 3.45, 2.0, 2.0, "页面状态", "Vue ref / Pinia")
flow_box(slide, 6.05, 2.0, 2.0, "API 封装", "Axios request.ts")
flow_box(slide, 8.65, 2.0, 2.0, "后端接口", "REST API")
flow_box(slide, 11.0, 2.0, 1.65, "页面更新", "渲染结果")
for i in range(4):
    arrow(slide, 2.85 + i * 2.6, 2.52, 3.45 + i * 2.6, 2.52)
card(slide, 0.8, 4.05, 3.75, 1.75, "登录状态", [
    "Token 存入 localStorage",
    "请求拦截器自动加 Authorization",
    "响应 401/403 时清理状态并跳转登录",
], BLUE)
card(slide, 4.8, 4.05, 3.75, 1.75, "开发联调", [
    "Vite 开发服务器负责前端热更新",
    "/api 请求代理到 Spring Boot 后端",
    "减少跨域问题，提高调试效率",
], CYAN)
card(slide, 8.8, 4.05, 3.75, 1.75, "权限路由", [
    "普通用户访问用户页面",
    "管理员登录后进入 /admin",
    "未登录访问受保护页面会跳转登录",
], GREEN)
footer(slide, page)
page += 1

slide = prs.slides.add_slide(prs.slide_layouts[6])
background(slide, prs)
title(slide, "项目管理做了什么", "把开发过程拆成阶段，并用文档和日志留下证据")
card(slide, 0.75, 1.35, 3.7, 4.95, "计划与分工", [
    "项目计划明确范围、目标、过程模型和进度安排",
    "三位成员按后端数据库、前端项目管理、APP报告分工",
    "开发中按模块推进，避免全部功能堆到最后联调",
], BLUE)
card(slide, 4.85, 1.35, 3.7, 4.95, "过程控制", [
    "需求获取和需求分析先确定功能边界",
    "系统设计和数据库设计再确定结构",
    "编码阶段按登录、硬件、推荐、评测、后台、APP 逐步完成",
    "build log 记录每日迭代和问题修复",
], CYAN)
card(slide, 8.95, 1.35, 3.7, 4.95, "验收准备", [
    "系统测试计划覆盖功能、性能、安全、兼容性",
    "结项报告总结完成情况、技术栈和不足",
    "答辩 PPT 按成员分工组织，便于现场回答",
], GREEN)
footer(slide, page)
page += 1

# APP/report section
section(prs, page, "APP 端 + 报告撰写", "负责 Android 封装、局域网访问、文档整理和工程证据链", ["移动端", "文档", "验收"])
page += 1

slide = prs.slides.add_slide(prs.slide_layouts[6])
background(slide, prs)
title(slide, "APP 端做了什么", "Android WebView 复用 Web 前端，让系统具备移动访问能力")
card(slide, 0.72, 1.45, 3.75, 4.95, "实现方式", [
    "使用 Kotlin + Android WebView 开发 APP 外壳",
    "将 Vue 构建产物复制到 assets/www",
    "APP 启动后加载本地 index.html",
    "不重复开发一套原生页面，减少工作量",
], BLUE)
card(slide, 4.82, 1.45, 3.75, 4.95, "通信逻辑", [
    "本地页面负责展示，后端数据仍来自 Spring Boot",
    "用户在设置页配置服务器 IP 和端口",
    "WebView 拦截 /api/* 请求并转发到局域网后端",
    "图片资源也从后端静态资源路径读取",
], CYAN)
card(slide, 8.92, 1.45, 3.75, 4.95, "为什么合理", [
    "Web 端功能已经完整，复用能保证双端一致",
    "课程项目重在完整流程和可运行系统",
    "WebView 性能足够支撑浏览、推荐、评测等场景",
    "局限：原生体验和离线能力仍可继续增强",
], GREEN)
footer(slide, page)
page += 1

slide = prs.slides.add_slide(prs.slide_layouts[6])
background(slide, prs)
title(slide, "报告撰写做了什么", "报告负责证明：不是只写代码，而是按软件工程方法完成系统")
card(slide, 0.75, 1.35, 3.7, 4.95, "前期文档", [
    "开题报告：说明背景、目标、意义和双端设想",
    "需求获取报告：明确用户角色、功能需求和约束",
    "需求分析报告：建立类图、行为模型和验收依据",
], BLUE)
card(slide, 4.85, 1.35, 3.7, 4.95, "设计文档", [
    "系统设计报告：说明页面、模块和事务逻辑",
    "数据库设计报告：说明表结构、命名规范和安全优化",
    "文档内容与代码模块对应，避免只停留在概念层",
], CYAN)
card(slide, 8.95, 1.35, 3.7, 4.95, "后期文档", [
    "系统测试计划：覆盖功能、健壮性、性能、安全和兼容性",
    "结项报告：总结完成情况、成果、测试和改进方向",
    "开发日志：补充说明每日迭代过程和问题解决记录",
], GREEN)
footer(slide, page)
page += 1

# Report evidence section
slide = prs.slides.add_slide(prs.slide_layouts[6])
background(slide, prs)
title(slide, "四份报告补充 1：项目计划支撑", "用项目计划说明范围、目标、工具、分工和进度")
card(slide, 0.68, 1.35, 3.75, 4.98, "项目范围与目标", [
    "范围：Web + Android 双端硬件推荐系统",
    "包含：用户认证、硬件数据库、排行算法、推荐算法、评测评论、收藏、管理员后台、Android APP",
    "不包含：在线支付、库存管理、物流跟踪",
    "目标：11 个功能页面、APK < 10MB、推荐计算 < 2s、API 响应 < 500ms、完成 8 份文档",
], BLUE)
card(slide, 4.8, 1.35, 3.75, 4.98, "过程模型与工具", [
    "采用 SPP 精简并行过程的迭代开发模型",
    "需求阶段：调研、竞品分析、需求报告",
    "设计阶段：架构设计、技术选型、数据库设计",
    "开发阶段：后端 API、前端页面、Android APP 分模块迭代",
    "工具：Git/GitHub、Spring Boot、Vue 3、MySQL、Postman、Android Studio、Draw.io",
], CYAN)
card(slide, 8.92, 1.35, 3.75, 4.98, "人员分工与进度", [
    "李锦轩：项目经理/后端开发/数据库设计，贡献 40%",
    "李俊泽：前端开发/测试/接口联调，贡献 30%",
    "黄国政：Android 开发/文档/系统测试，贡献 30%",
    "进度按需求、设计、开发、测试、交付推进，最终在 2026 年 5 月完成验收材料",
], GREEN)
footer(slide, page)
page += 1

slide = prs.slides.add_slide(prs.slide_layouts[6])
background(slide, prs)
title(slide, "四份报告补充 2：数据库设计支撑", "数据库报告把系统对象、表结构、规范和安全策略讲清楚")
card(slide, 0.68, 1.35, 3.75, 4.98, "数据库环境", [
    "数据库系统：MySQL 8.0",
    "数据库名称：pcadvisor",
    "字符集：utf8mb4，支持中文和 Emoji",
    "地址：localhost:3306",
    "ORM：MyBatis-Plus 3.5.5 + Spring Data JPA",
    "引擎：InnoDB，支持事务和约束",
], BLUE)
card(slide, 4.8, 1.35, 3.75, 4.98, "核心实体与关系", [
    "用户体系：sys_user、OrdinaryUser、Administrator",
    "硬件体系：Hardware_Inf、CPU_Inf、GraphicsCard_Inf、Motherboard_Inf",
    "UGC 体系：Evaluation、Comment、Collect、Media_Inf",
    "关键关系：用户与身份一对一，硬件主表与类型详情表一对一，硬件与评测一对多",
], CYAN)
card(slide, 8.92, 1.35, 3.75, 4.98, "规范与安全", [
    "命名规范：idx_ 表示索引，uk_ 表示唯一约束",
    "主表字段使用下划线，用户表兼容 camelCase",
    "应用层隔离：用户不能直接操作数据库",
    "参数化查询和 ORM 映射降低 SQL 注入风险",
    "密码字段使用加密存储，角色字段区分 ADMIN/ORDINARY",
], GREEN)
footer(slide, page)
page += 1

slide = prs.slides.add_slide(prs.slide_layouts[6])
background(slide, prs)
title(slide, "四份报告补充 3：系统测试计划支撑", "测试计划说明系统如何证明“能正常运行”")
card(slide, 0.68, 1.35, 3.75, 4.98, "测试范围", [
    "功能测试：注册登录、Token、硬件浏览、排行、推荐、评测、评论、收藏、管理员 CRUD、审核、上传、APP 代理",
    "健壮性测试：空值、超长文本、非法 Token、过期 Token、大文件上传拒绝",
    "UI 测试：Web 响应式、Android WebView、弹窗、图表、图片灯箱、后台主题",
], BLUE)
card(slide, 4.8, 1.35, 3.75, 4.98, "测试方法与环境", [
    "黑盒测试：基于需求规格覆盖功能点",
    "白盒测试：针对推荐算法、JWT、密码加密",
    "集成测试：验证 Controller 完整请求响应链路",
    "环境：Windows 11、JDK 21、Spring Boot 3.2.2、MySQL 8.0.36、Vue 3、Node.js 22、Android SDK 35",
], CYAN)
card(slide, 8.92, 1.35, 3.75, 4.98, "完成准则", [
    "核心功能测试通过率达到 100%",
    "非功能测试通过率不低于 95%",
    "无 P0/P1 级缺陷，不出现系统崩溃、数据丢失或严重安全漏洞",
    "回归测试通过，已修复问题不再次出现",
    "性能目标：API < 500ms，推荐 < 2s，5MB 内图片上传 < 3s",
], GREEN)
footer(slide, page)
page += 1

slide = prs.slides.add_slide(prs.slide_layouts[6])
background(slide, prs)
title(slide, "四份报告补充 4：结项报告支撑", "结项报告证明最终完成了哪些成果，以及有哪些经验教训")
card(slide, 0.68, 1.35, 3.75, 4.98, "完成成果", [
    "需求分析：整理功能需求 22 项",
    "数据库：完成 11 张表，采用主表 + 详情子表架构",
    "后端：9 个 Controller，30+ 个 API 接口",
    "前端：11 个 Vue 页面组件，支持响应式布局",
    "Android：完成 WebView APP，APK 约 6.9MB",
    "文档：完成 8 份规范技术文档",
], BLUE)
card(slide, 4.8, 1.35, 3.75, 4.98, "质量评价", [
    "正确性：核心功能经人工测试符合预期",
    "健壮性：全局异常处理覆盖认证、授权、业务和资源场景",
    "安全性：BCrypt、JWT、Spring Security、参数化 SQL",
    "扩展性：硬件类型通过枚举和详情表扩展",
    "兼容性：Chrome/Firefox/Edge 与 Android WebView API 24+",
], CYAN)
card(slide, 8.92, 1.35, 3.75, 4.98, "经验教训", [
    "技术选型要提前验证兼容性，尤其是 WebView 与 ES Module",
    "前后端分离需要尽早统一错误码和响应格式",
    "文件上传要注意 Windows/Linux 框架行为差异",
    "数据库主表 + 子表模式比单一宽表更适合硬件领域",
    "文档要与开发同步，避免后期集中赶工",
], GREEN)
footer(slide, page)
page += 1

# Questions
slide = prs.slides.add_slide(prs.slide_layouts[6])
background(slide, prs)
title(slide, "自设答辩问题 1：后端与数据库", "问题要答到“原因 + 实现 + 效果”")
card(slide, 0.72, 1.35, 5.75, 2.15, "Q1：为什么数据库采用“硬件主表 + 类型详情表”？", [
    "A：硬件都有名称、品牌、型号、价格等公共字段，但 CPU、显卡、主板的参数完全不同。公共字段放 Hardware_Inf，专属字段放 CPU_Inf、GraphicsCard_Inf、Motherboard_Inf，可以减少空字段和冗余，也方便后续扩展更多硬件类型。",
], BLUE, 12.8, 10.2)
card(slide, 6.85, 1.35, 5.75, 2.15, "Q2：推荐算法的逻辑是什么？", [
    "A：先从审核通过的硬件中取候选，再用预算、用途、兼容性等硬性条件过滤，最后按品牌、系列、板型等软性偏好打分排序，返回匹配度最高的结果。这样既保证推荐可用，又体现个性化。",
], CYAN, 12.8, 10.2)
card(slide, 0.72, 4.15, 5.75, 2.15, "Q3：后端如何保证权限安全？", [
    "A：登录成功后生成 JWT，前端请求时自动携带；后端过滤器解析 Token 并识别用户身份。普通用户访问个人中心和评测功能，管理员才能访问 /api/admin 下的管理接口。",
], GREEN, 12.8, 10.2)
card(slide, 6.85, 4.15, 5.75, 2.15, "Q4：数据库实际部署在哪里？", [
    "A：项目连接本机 MySQL，配置地址是 jdbc:mysql://localhost:3306/pcadvisor。实际数据目录在 D:\\MySQL\\MySQL Server 8.0\\Data\\pcadvisor。",
], BLUE, 12.8, 10.2)
footer(slide, page)
page += 1

slide = prs.slides.add_slide(prs.slide_layouts[6])
background(slide, prs)
title(slide, "自设答辩问题 2：前端与项目管理", "突出用户体验、联调方式和进度控制")
card(slide, 0.72, 1.35, 5.75, 2.15, "Q1：前端如何和后端交互？", [
    "A：前端通过 Axios 请求 REST API。request.ts 统一处理 baseURL、Token、错误响应；开发环境中 Vite 把 /api 代理到 Spring Boot，解决跨域和联调问题。",
], BLUE, 12.8, 10.2)
card(slide, 6.85, 1.35, 5.75, 2.15, "Q2：用户和管理员页面如何区分？", [
    "A：Vue Router 根据路由和角色控制访问。普通用户进入首页、硬件、排行、推荐、评测、个人中心；管理员登录后进入 /admin 后台，进行硬件管理和内容审核。",
], CYAN, 12.8, 10.2)
card(slide, 0.72, 4.15, 5.75, 2.15, "Q3：项目进度如何体现？", [
    "A：项目计划中有阶段划分，build log 记录每日开发进展；结项报告总结实际完成情况。实际流程基本按需求、设计、编码、测试、交付推进。",
], GREEN, 12.8, 10.2)
card(slide, 6.85, 4.15, 5.75, 2.15, "Q4：前端最能现场展示什么？", [
    "A：可以展示登录注册、硬件筛选和详情、排行页面、推荐页面、发布评测、个人收藏，以及管理员后台的硬件管理和评测审核。",
], BLUE, 12.8, 10.2)
footer(slide, page)
page += 1

slide = prs.slides.add_slide(prs.slide_layouts[6])
background(slide, prs)
title(slide, "自设答辩问题 3：APP 与报告", "说明选择依据、局限性和文档质量")
card(slide, 0.72, 1.35, 5.75, 2.15, "Q1：为什么 APP 端选择 WebView？", [
    "A：因为 Web 端功能已经完整，WebView 可以复用页面，保证双端一致并减少重复开发。对于硬件浏览、推荐、评测这类信息系统，WebView 的性能和体验基本满足课程项目需求。",
], BLUE, 12.8, 10.2)
card(slide, 6.85, 1.35, 5.75, 2.15, "Q2：APP 如何连接后端？", [
    "A：APP 加载本地前端资源，数据请求通过 /api/* 转发到局域网中的 Spring Boot 后端。用户可以在设置页配置服务器 IP 和端口，并进行连接测试。",
], CYAN, 12.8, 10.2)
card(slide, 0.72, 4.15, 5.75, 2.15, "Q3：报告如何体现软件工程过程？", [
    "A：文档覆盖开题、需求获取、需求分析、系统设计、数据库设计、测试计划和结项报告，能从问题提出追溯到设计、实现和测试验证。",
], GREEN, 12.8, 10.2)
card(slide, 6.85, 4.15, 5.75, 2.15, "Q4：项目还有哪些不足？", [
    "A：推荐算法目前主要基于规则和静态数据，后续可以接入实时价格和真实性能跑分；APP 端也可以增加离线缓存、通知和扫码比价等原生能力。",
], BLUE, 12.8, 10.2)
footer(slide, page)
page += 1

# Goals
slide = prs.slides.add_slide(prs.slide_layouts[6])
background(slide, prs)
title(slide, "课程目标对应说明", "把验收目标和工程证据一一对应")
card(slide, 0.72, 1.25, 5.75, 1.4, "目标 1：设计满足需求的解决方案", [
    "体现：围绕硬件选购难题设计硬件库、排行、推荐、评测社区和审核后台；结合硬件领域参数设计数据模型。",
], BLUE, 12.5, 9.8)
card(slide, 6.85, 1.25, 5.75, 1.4, "目标 2：完成软件工程全过程", [
    "体现：需求获取、需求分析、系统设计、数据库设计、编码实现、测试计划和结项报告完整存在。",
], CYAN, 12.5, 9.8)
card(slide, 0.72, 3.25, 5.75, 1.4, "目标 3：使用现代工具和平台", [
    "体现：Spring Boot、MySQL、MyBatis-Plus、Vue 3、TypeScript、Vite、Android WebView 等技术协同使用。",
], GREEN, 12.5, 9.8)
card(slide, 6.85, 3.25, 5.75, 1.4, "目标 4：小组协作完成开发", [
    "体现：三位成员分工明确，开发日志和结项报告记录了阶段成果与协作过程。",
], BLUE, 12.5, 9.8)
text(slide, 0.92, 5.8, 11.3, 0.5, "答辩主线：我们不是只做了页面，而是从需求、设计、数据库、后端、前端、APP、测试和文档交付完整走了一遍软件开发过程。", 14, True, BLUE_DARK, PP_ALIGN.CENTER)
footer(slide, page)
page += 1

# End
slide = prs.slides.add_slide(prs.slide_layouts[6])
background(slide, prs)
panel = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(1.15), Inches(11.35), Inches(5.4))
panel.fill.solid()
panel.fill.fore_color.rgb = WHITE
panel.line.color.rgb = LINE
text(slide, 1.65, 1.85, 10.0, 0.52, "总结", 29, True, BLUE_DARK, PP_ALIGN.CENTER)
text(slide, 1.95, 2.7, 9.4, 1.25, "本系统以电脑硬件选购为具体场景，完成了 Web 端、后端、数据库、Android APP 和软件工程文档。系统逻辑清晰：前端负责交互，后端负责业务，数据库负责持久化，APP 复用 Web 成果，报告负责证明开发过程。", 16, False, TEXT, PP_ALIGN.CENTER)
for i, label in enumerate(["功能闭环", "技术完整", "分工明确", "文档可追溯"]):
    pill(slide, 2.1 + i * 2.3, 4.55, 1.75, label, [BLUE, CYAN, GREEN, BLUE_DARK][i])
text(slide, 4.6, 6.05, 4.1, 0.35, "谢谢老师，请批评指正", 18, True, BLUE_DARK, PP_ALIGN.CENTER)
footer(slide, page)

OUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUT)
print(OUT.resolve())
